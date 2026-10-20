"""Utility script to generate sample PDF and DOCX resume files for testing."""

import os
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
import docx


def generate_sample_pdf(txt_path: str, pdf_path: str):
    """Generate a clean PDF resume from text file."""
    with open(txt_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = SimpleDocTemplate(
        pdf_path,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'ResumeTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=6
    )
    heading_style = ParagraphStyle(
        'ResumeHeading',
        parent=styles['Heading2'],
        fontSize=12,
        leading=16,
        textColor=colors.HexColor("#1e3a8a"),
        spaceBefore=10,
        spaceAfter=4,
        borderPadding=(0, 0, 2, 0)
    )
    body_style = ParagraphStyle(
        'ResumeBody',
        parent=styles['Normal'],
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#334155")
    )

    story = []
    lines = content.split("\n")
    is_first_line = True

    for line in lines:
        stripped = line.strip()
        if not stripped:
            story.append(Spacer(1, 4))
            continue

        if is_first_line:
            story.append(Paragraph(stripped, title_style))
            is_first_line = False
        elif stripped.isupper() and len(stripped) < 40:
            story.append(Paragraph(f"<b>{stripped}</b>", heading_style))
        else:
            safe_text = stripped.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(safe_text, body_style))

    doc.build(story)
    print(f"[SUCCESS] Generated PDF: {pdf_path}")


def generate_sample_docx(txt_path: str, docx_path: str):
    """Generate a clean DOCX resume from text file."""
    with open(txt_path, "r", encoding="utf-8") as f:
        content = f.read()

    doc = docx.Document()
    lines = content.split("\n")
    is_first_line = True

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        if is_first_line:
            p = doc.add_heading(stripped, level=1)
            is_first_line = False
        elif stripped.isupper() and len(stripped) < 40:
            doc.add_heading(stripped, level=2)
        elif stripped.startswith("- "):
            doc.add_paragraph(stripped[2:], style='List Bullet')
        else:
            doc.add_paragraph(stripped)

    doc.save(docx_path)
    print(f"[SUCCESS] Generated DOCX: {docx_path}")


if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    resumes_dir = os.path.join(base_dir, "data", "sample_resumes")

    alex_txt = os.path.join(resumes_dir, "alex_chen_data_scientist.txt")
    alex_pdf = os.path.join(resumes_dir, "alex_chen_data_scientist.pdf")

    priya_txt = os.path.join(resumes_dir, "priya_patel_fullstack_dev.txt")
    priya_docx = os.path.join(resumes_dir, "priya_patel_fullstack_dev.docx")

    if os.path.exists(alex_txt):
        generate_sample_pdf(alex_txt, alex_pdf)

    if os.path.exists(priya_txt):
        generate_sample_docx(priya_txt, priya_docx)
