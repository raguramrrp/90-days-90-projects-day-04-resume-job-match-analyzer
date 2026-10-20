"""Document Parser and Text Extraction Engine.

Handles text extraction from PDF, DOCX, and TXT files with encoding fallbacks,
noise reduction, contact extraction, and section identification.
"""

from __future__ import annotations

import io
import re
from typing import Any, Dict, Optional
import pypdf
import docx


class DocumentParseError(Exception):
    """Raised when document parsing fails."""
    pass


def clean_text(raw_text: str) -> str:
    """Normalize extracted text by removing control characters and redundant spaces."""
    if not raw_text:
        return ""
    # Replace non-breaking spaces and carriage returns
    text = raw_text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    # Replace multiple spaces/tabs with single space
    text = re.sub(r"[ \t]+", " ", text)
    # Replace more than 2 consecutive newlines with 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extract text from PDF byte content using pypdf."""
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        for page_idx, page in enumerate(reader.pages):
            page_text = page.extract_text() or ""
            if page_text.strip():
                pages_text.append(page_text)
        extracted = "\n".join(pages_text)
        if not extracted.strip():
            raise DocumentParseError(
                "PDF extracted text was empty. The file may be image-only (scanned) or password-protected."
            )
        return clean_text(extracted)
    except Exception as e:
        if isinstance(e, DocumentParseError):
            raise
        raise DocumentParseError(f"Failed to parse PDF document: {e}") from e


def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extract text from DOCX byte content using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text.strip())
                        
        extracted = "\n".join(paragraphs)
        if not extracted.strip():
            raise DocumentParseError("DOCX extracted text was empty.")
        return clean_text(extracted)
    except Exception as e:
        if isinstance(e, DocumentParseError):
            raise
        raise DocumentParseError(f"Failed to parse DOCX document: {e}") from e


def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extract text from raw bytes with encoding fallbacks."""
    encodings = ["utf-8", "latin-1", "cp1252", "utf-16"]
    for enc in encodings:
        try:
            decoded = file_bytes.decode(enc)
            return clean_text(decoded)
        except (UnicodeDecodeError, LookupError):
            continue
    raise DocumentParseError("Unable to decode TXT file with supported encodings.")


def extract_text_from_file(file_bytes: bytes, filename: str) -> str:
    """Extract and clean text from uploaded file based on its extension."""
    if not file_bytes:
        raise DocumentParseError("Uploaded file is empty (0 bytes).")
    
    fname = filename.lower()
    if fname.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif fname.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif fname.endswith(".txt") or fname.endswith(".md"):
        return extract_text_from_txt(file_bytes)
    else:
        raise DocumentParseError(
            f"Unsupported file format '{filename}'. Please upload a PDF (.pdf), Word Document (.docx), or Text file (.txt)."
        )


def extract_contact_info(text: str) -> Dict[str, Optional[str]]:
    """Extract contact information (email, phone, LinkedIn, GitHub, Website) from resume."""
    contacts: Dict[str, Optional[str]] = {
        "email": None,
        "phone": None,
        "linkedin": None,
        "github": None,
        "website": None
    }
    
    # Email regex
    email_match = re.search(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text)
    if email_match:
        contacts["email"] = email_match.group(0).rstrip(".")
        
    # Phone regex (international & US formats)
    phone_match = re.search(
        r"(?:(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{2,4}\)?[-.\s]?)?\d{3,4}[-.\s]?\d{4})", 
        text
    )
    if phone_match:
        phone_candidate = phone_match.group(0).strip()
        # Ensure it has at least 7 digits to avoid matching random years/numbers
        if sum(c.isdigit() for c in phone_candidate) >= 7:
            contacts["phone"] = phone_candidate
            
    # LinkedIn
    linkedin_match = re.search(r"(?:https?:\/\/)?(?:www\.)?linkedin\.com\/in\/[a-zA-Z0-9_-]+", text, re.IGNORECASE)
    if linkedin_match:
        contacts["linkedin"] = linkedin_match.group(0)
        
    # GitHub
    github_match = re.search(r"(?:https?:\/\/)?(?:www\.)?github\.com\/[a-zA-Z0-9_-]+", text, re.IGNORECASE)
    if github_match:
        contacts["github"] = github_match.group(0)
        
    # General portfolio / website (if distinct from github / linkedin / email domain)
    portfolio_matches = re.finditer(
        r"(?<!@)\b(?:https?:\/\/)?(?:www\.)?([a-zA-Z0-9-]+(?:\.[a-zA-Z0-9-]+)*\.(?:io|dev|me|tech|com|org|net)(?:\/[^\s]*)?)", 
        text, 
        re.IGNORECASE
    )
    for m in portfolio_matches:
        url = m.group(0).strip().rstrip(".,;)")
        url_lower = url.lower()
        if "linkedin.com" not in url_lower and "github.com" not in url_lower:
            if contacts["email"] and url_lower in contacts["email"].lower():
                continue
            contacts["website"] = url
            break
            
    return contacts


def extract_sections(text: str) -> Dict[str, str]:
    """Segment resume text into standard thematic sections."""
    section_patterns = {
        "Summary": r"(?:professional\s+summary|profile|executive\s+summary|objective|about\s+me)",
        "Experience": r"(?:work\s+experience|professional\s+experience|experience|employment\s+history|internships)",
        "Education": r"(?:education|academic\s+background|qualifications|academic\s+credentials)",
        "Skills": r"(?:technical\s+skills|skills\s*(?:&|and)?\s*competencies|core\s+competencies|technologies|skills)",
        "Projects": r"(?:projects|academic\s+projects|personal\s+projects|key\s+projects)",
        "Certifications": r"(?:certifications|licenses\s*(?:&|and)?\s*certifications|professional\s+certifications|awards)"
    }
    
    # Find all headers and their positions
    matches = []
    lines = text.split("\n")
    current_char_idx = 0
    
    for line in lines:
        stripped = line.strip()
        # Section headers are typically short (1-5 words)
        if 2 <= len(stripped) <= 45:
            for sec_name, pattern in section_patterns.items():
                if re.fullmatch(rf"(?i)(?:\d+[\.\)]\s*)?{pattern}(?:\s*:)?", stripped):
                    matches.append((current_char_idx, sec_name, stripped))
                    break
        current_char_idx += len(line) + 1  # newline length
        
    sections: Dict[str, str] = {}
    if not matches:
        sections["Full Text"] = text
        return sections
        
    matches.sort(key=lambda x: x[0])
    
    # Extract text between matches
    for idx, (start_idx, sec_name, header) in enumerate(matches):
        content_start = start_idx + len(header)
        if idx + 1 < len(matches):
            content_end = matches[idx + 1][0]
            sec_content = text[content_start:content_end]
        else:
            sec_content = text[content_start:]
        sections[sec_name] = clean_text(sec_content)
        
    return sections
