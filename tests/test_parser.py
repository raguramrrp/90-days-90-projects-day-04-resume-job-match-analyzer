"""Tests for document parsing and text extraction."""

import io
import pytest
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph
from reportlab.lib.styles import getSampleStyleSheet

from src.parser import (
    clean_text,
    extract_text_from_file,
    extract_text_from_pdf,
    extract_text_from_docx,
    extract_text_from_txt,
    extract_contact_info,
    extract_sections,
    DocumentParseError
)


def test_clean_text():
    dirty = "Hello   world!\r\n\r\n\n\nThis is\xa0a test.   "
    cleaned = clean_text(dirty)
    assert cleaned == "Hello world!\n\nThis is a test."


def test_extract_text_from_txt():
    content = "Software Engineer Resume\nPython, Django, AWS"
    raw_bytes = content.encode("utf-8")
    extracted = extract_text_from_txt(raw_bytes)
    assert "Software Engineer" in extracted
    assert "Python, Django, AWS" in extracted


def test_extract_text_from_docx():
    doc = docx.Document()
    doc.add_heading("Jane Doe Resume", level=1)
    doc.add_paragraph("Full-Stack Developer with React and Node.js.")
    
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    extracted = extract_text_from_docx(docx_bytes)
    assert "Jane Doe Resume" in extracted
    assert "Full-Stack Developer" in extracted


def test_extract_text_from_pdf():
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("John Smith", styles['Heading1']),
        Paragraph("AI Engineer specializing in NLP and PyTorch", styles['Normal'])
    ]
    doc.build(story)
    pdf_bytes = buf.getvalue()

    extracted = extract_text_from_pdf(pdf_bytes)
    assert "John Smith" in extracted
    assert "AI Engineer" in extracted


def test_extract_text_from_file_dispatcher():
    txt_bytes = b"Sample resume text"
    assert extract_text_from_file(txt_bytes, "test.txt") == "Sample resume text"

    with pytest.raises(DocumentParseError, match="Unsupported file format"):
        extract_text_from_file(txt_bytes, "test.exe")

    with pytest.raises(DocumentParseError, match="Uploaded file is empty"):
        extract_text_from_file(b"", "test.pdf")


def test_extract_contact_info():
    sample = (
        "John Doe\n"
        "Email: john.doe@example.com\n"
        "Phone: +1 (555) 345-6789\n"
        "LinkedIn: https://linkedin.com/in/johndoe\n"
        "GitHub: https://github.com/johndoe\n"
        "Portfolio: https://johndoe.dev\n"
    )
    contacts = extract_contact_info(sample)
    assert contacts["email"] == "john.doe@example.com"
    assert "555" in contacts["phone"]
    assert "linkedin.com/in/johndoe" in contacts["linkedin"]
    assert "github.com/johndoe" in contacts["github"]
    assert "johndoe.dev" in contacts["website"]


def test_extract_sections():
    sample = (
        "ALEX CHEN\n\n"
        "PROFESSIONAL SUMMARY\n"
        "Experienced Data Scientist with 5 years in AI.\n\n"
        "TECHNICAL SKILLS\n"
        "Python, SQL, PyTorch, Docker\n\n"
        "PROFESSIONAL EXPERIENCE\n"
        "Senior ML Engineer at TechCorp.\n\n"
        "EDUCATION\n"
        "BS in Computer Science\n"
    )
    sections = extract_sections(sample)
    assert "Summary" in sections
    assert "Skills" in sections
    assert "Experience" in sections
    assert "Education" in sections
    assert "Experienced Data Scientist" in sections["Summary"]
    assert "Python, SQL" in sections["Skills"]
